from types import SimpleNamespace

import numpy as np
import pytest
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QApplication, QMessageBox

import bictool.application.worker as worker_module
import bictool.presentation.views.main_window as main_window_module
from bictool.integration.dicom_file_loader import DicomVolume, SeriesCandidate
from bictool.presentation.viewmodels.roi_selection import RoiRect
from bictool.presentation.viewmodels.windowing import WINDOW_PRESETS
from bictool.presentation.views.main_window import MainWindowView


def _fake_series_candidate(label: str, num_slices: int = 1, pixel_spacing: str = "0.5") -> SeriesCandidate:
    reference = SimpleNamespace(
        PixelSpacing=[pixel_spacing, pixel_spacing],
        SliceThickness=None,
        PatientID="ANON-0001",
        PatientName="ANON^Patient",
        PatientSex="O",
        StudyID="STU-0001",
        StudyDate="20260101",
        StudyDescription="CT Schadel Spirale",
        SeriesDescription="Head 0.5 Vol",
        SeriesNumber="4",
        InstitutionName="Test Hospital",
        Manufacturer="TestCorp",
        ManufacturerModelName="Scanner X",
        AcquisitionDate="20260101",
        AcquisitionTime="104500",
        KVP=120.0,
        XRayTubeCurrent=170,
        Modality="CT",
        CTDIvol=None,
        DLP=None,
        ConvolutionKernel=None,
    )
    return SeriesCandidate(
        label=label,
        sorted_paths=[f"fake_{label}_{i}.dcm" for i in range(num_slices)],
        slice_locations_mm=[float(i) for i in range(num_slices)],
        _reference_header=reference,
    )


def _fake_dicom_volume(num_slices: int = 1) -> DicomVolume:
    # slice_paths are never actually read - the pixel cache is pre-populated
    # below, so DicomVolume.get_slice_hu never touches disk for these fakes.
    volume = DicomVolume(
        slice_paths=[f"fake_slice_{i}.dcm" for i in range(num_slices)],
        pixel_spacing_mm=0.5,
        slice_thickness_mm=1.0,
        slice_locations_mm=[float(i) for i in range(num_slices)],
        patient_id="ANON-0001",
        patient_name="ANON Patient",
        patient_sex="O",
        study_id="STU-0001",
        study_date="20260101",
        study_description="CT Schadel Spirale",
        series_description="Head 0.5 Vol",
        series_number="4",
        institution_name="Test Hospital",
        manufacturer_model="TestCorp Scanner X",
        acquisition_date="20260101",
        acquisition_time="10:45:00",
        kvp=120.0,
        tube_current_ma=170.0,
        transfer_syntax_name="Implicit VR Little Endian",
        modality="CT",
        ctdi_vol_mgy=None,
        dlp_mgy_cm=None,
        reconstruction=None,
    )
    volume._slice_cache = {i: np.zeros((60, 60), dtype=np.float64) for i in range(num_slices)}
    return volume


def _load_fake_volume(window: MainWindowView, num_slices: int = 1) -> DicomVolume:
    """Mirrors what `_on_volume_loaded` does, without going through QFileDialog/DicomLoadWorker."""
    volume = _fake_dicom_volume(num_slices)
    window._on_volume_loaded(volume)
    return volume


def _set_roi(window: MainWindowView, key: str, rect: RoiRect) -> None:
    """Mirrors drawing `rect` on the image and clicking "Set ROI1"/"Set
    ROI2" (`_save_draft_as`), without driving real mouse events -
    `RoiViewport`'s draw-then-assign model means a rectangle must exist as
    its current draft before a Set ROI button can save it.
    """
    window._roi_viewport._current_rect = rect
    # Plain "roi1"/"roi2" = the falx pair (the app's slots are "falx.roi1", ...).
    window._save_draft_as(key if "." in key else f"falx.{key}")


def _fake_result() -> SimpleNamespace:
    return SimpleNamespace(
        mtf_area=0.418,
        mtf_xaxis=np.array([0.0, 0.5, 1.0]),
        mtf=np.array([[1.0, 1.0], [0.6, 0.5], [0.2, 0.1]]),
        nps_roi1=SimpleNamespace(
            var_nps=0.00187,
            f_avg=np.array([0.0, 0.5, 1.0]),
            nps_measured_radagv=np.array([0.002, 0.0015, 0.0005]),
        ),
        nps_roi2=SimpleNamespace(
            var_nps=0.00291,
            f_avg=np.array([0.0, 0.5, 1.0]),
            nps_measured_radagv=np.array([0.003, 0.0022, 0.0009]),
        ),
        ratio=SimpleNamespace(
            ratio_lf_hf=1.86,
            f_int=np.array([0.0, 0.5, 1.0]),
            ratio_nps_mtf=np.array([2.0, 1.5, 1.0]),
        ),
        contrast_roi1=812.3,
        predicted_label=1,
        decision_value=1.42,
        feature_vector=[0.418, 0.00142, 0.00187, 12.4, 1.86, -0.031, 0.34],
    )


def test_window_has_expected_title_and_menus():
    window = MainWindowView()
    assert window.windowTitle() == "BICTool"
    menu_titles = [action.text().replace("&", "") for action in window.menuBar().actions()]
    assert menu_titles == ["File", "Study", "View", "Tools", "Settings", "Export"]


def test_metric_status_starts_idle_not_queued():
    # "queued" implies something is waiting to run, which isn't true before
    # Run All has ever been clicked - that misleading initial state was the
    # bug being fixed here.
    window = MainWindowView()
    for key in ("mtf", "nps", "ratio", "contrast"):
        assert window._metric_status_items[key].text() == "—"


def test_series_combo_starts_disabled_with_no_items():
    window = MainWindowView()
    assert window._series_combo.count() == 0
    assert window._series_combo.isEnabled() is False


def test_single_series_discovered_auto_loads_and_combo_stays_disabled(monkeypatch):
    # These candidates' sorted_paths aren't real files - only load_series's
    # metadata handling is under test here, not pixel decoding, so stub out
    # the one call that would otherwise touch disk.
    monkeypatch.setattr(DicomVolume, "get_slice_hu", lambda self, index: np.zeros((8, 8), dtype=np.float64))
    window = MainWindowView()
    candidate = _fake_series_candidate("CT · 3 slice(s) · 8×8", num_slices=3)

    window._on_series_discovered([candidate])

    assert window._series_combo.count() == 1
    assert window._series_combo.currentText() == "CT · 3 slice(s) · 8×8"
    # One candidate is nothing to choose between - shown, but not interactive.
    assert window._series_combo.isEnabled() is False
    assert window.view_model.study.volume.num_slices == 3


def test_multiple_series_discovered_populates_combo_and_enables_it(monkeypatch):
    monkeypatch.setattr(DicomVolume, "get_slice_hu", lambda self, index: np.zeros((8, 8), dtype=np.float64))
    window = MainWindowView()
    ct_candidate = _fake_series_candidate("CT · 3 slice(s) · 8×8", num_slices=3)
    scout_candidate = _fake_series_candidate("OT · 1 slice(s) · 8×8", num_slices=1)

    window._on_series_discovered([ct_candidate, scout_candidate])

    assert window._series_combo.count() == 2
    assert window._series_combo.isEnabled() is True
    # Best-first ordering (already sorted by discover_series) - the first
    # candidate is auto-loaded.
    assert window._series_combo.currentIndex() == 0
    assert window.view_model.study.volume.num_slices == 3


def test_switching_series_combo_selection_reloads_the_volume(monkeypatch):
    monkeypatch.setattr(DicomVolume, "get_slice_hu", lambda self, index: np.zeros((8, 8), dtype=np.float64))
    window = MainWindowView()
    ct_candidate = _fake_series_candidate("CT · 3 slice(s) · 8×8", num_slices=3)
    scout_candidate = _fake_series_candidate("OT · 1 slice(s) · 8×8", num_slices=1)
    window._on_series_discovered([ct_candidate, scout_candidate])
    assert window.view_model.study.volume.num_slices == 3

    window._series_combo.setCurrentIndex(1)

    assert window.view_model.study.volume.num_slices == 1


def _overlay_text(window: MainWindowView, corner: str) -> str:
    return window._roi_viewport._overlay_labels[corner].text()


def test_sidebar_no_longer_shows_per_field_study_info_labels():
    # Patient/study/technical identity moved onto the image itself as corner
    # overlay text - the sidebar should only have the series picker left.
    window = MainWindowView()
    assert not hasattr(window, "_study_info_labels")


def test_loading_a_volume_shows_patient_and_technical_details_as_corner_overlay_text():
    # Field placement mirrors the researched RadiAnt/MicroDicom convention:
    # identity (name/ID/sex/descriptions) top-left, institution/scanner/
    # acquisition time top-right, geometry/series-and-image-count bottom-left,
    # exposure/display state bottom-right.
    window = MainWindowView()
    volume = _load_fake_volume(window, num_slices=3)

    top_left = _overlay_text(window, "top_left")
    assert volume.patient_name in top_left
    assert "ANON-0001" in top_left
    assert volume.study_description in top_left
    assert volume.series_description in top_left

    top_right = _overlay_text(window, "top_right")
    assert volume.institution_name in top_right
    assert volume.manufacturer_model in top_right

    bottom_left = _overlay_text(window, "bottom_left")
    assert f"{volume.pixel_spacing_mm:.3f}" in bottom_left
    assert "Image 1/3" in bottom_left

    bottom_right = _overlay_text(window, "bottom_right")
    assert f"{volume.kvp:.0f} kV" in bottom_right
    assert f"{volume.tube_current_ma:.0f} mA" in bottom_right


def test_changing_slice_updates_the_overlays_slice_number():
    window = MainWindowView()
    _load_fake_volume(window, num_slices=3)

    window._show_slice(1)

    assert "Image 2/3" in _overlay_text(window, "bottom_left")


def test_switching_series_updates_overlay_to_that_series_own_pixel_spacing(monkeypatch):
    # Different series in the same study can legitimately have different
    # pixel spacing (e.g. a different reconstruction) - the overlay must
    # reflect whichever series is actually loaded, not a stale reading.
    monkeypatch.setattr(DicomVolume, "get_slice_hu", lambda self, index: np.zeros((8, 8), dtype=np.float64))
    window = MainWindowView()
    fine_candidate = _fake_series_candidate("CT · 3 slice(s) · 8×8", num_slices=3, pixel_spacing="0.488")
    coarse_candidate = _fake_series_candidate("CT · 1 slice(s) · 8×8", num_slices=1, pixel_spacing="0.700")
    window._on_series_discovered([fine_candidate, coarse_candidate])
    assert "0.488" in _overlay_text(window, "bottom_left")

    window._series_combo.setCurrentIndex(1)

    assert "0.700" in _overlay_text(window, "bottom_left")
    assert "0.488" not in _overlay_text(window, "bottom_left")


def test_window_preset_change_updates_the_overlays_wl_ww():
    window = MainWindowView()
    _load_fake_volume(window, num_slices=1)
    bone_width, bone_center = WINDOW_PRESETS["Bone"]

    window._on_window_preset_clicked("Bone")

    assert f"WL: {bone_center:.0f}" in _overlay_text(window, "bottom_right")
    assert f"WW: {bone_width:.0f}" in _overlay_text(window, "bottom_right")


def test_zoom_change_updates_the_overlays_zoom_percent():
    # `zoom_changed`'s handler re-reads the live zoom via
    # `current_zoom_percent()` rather than trusting the signal's payload
    # (it's just a "something changed, refresh" nudge) - so this drives a
    # real scale change first, matching what `wheelEvent` does internally,
    # instead of asserting against a value the handler wouldn't actually use.
    window = MainWindowView()
    _load_fake_volume(window, num_slices=1)

    window._roi_viewport.scale(2.0, 2.0)
    window._roi_viewport.zoom_changed.emit(window._roi_viewport.current_zoom_percent())

    expected = f"Zoom: {window._roi_viewport.current_zoom_percent():.0f}%"
    assert expected in _overlay_text(window, "bottom_right")


def test_export_action_disabled_until_a_result_exists():
    # Export (FR-07) is real now - disabled (not a silent no-op) only until
    # Run All has actually produced something to export.
    window = MainWindowView()
    assert window._export_action.isEnabled() is False
    assert "Run metrics first" in window._export_action.toolTip()


def test_trace_panel_receives_log_messages():
    import logging

    window = MainWindowView()
    window._show_trace_action.setChecked(True)  # off by default (Settings > Show Trace Log)

    logging.getLogger("bictool.some.module").info("a distinctive test message")

    assert "a distinctive test message" in window._trace_log.toPlainText()


def test_run_all_seeds_the_mtf_active_contour_on_roi1s_bright_line(monkeypatch):
    # calc_mtf's active contour must start inside the edge structure - a
    # manual ROI1 (no segmentation) is seeded on its brightest central line.
    captured = {}

    def fake_evaluate(
        img_roi1, img_roi2, pix_spc, act_cont_mask_x, act_cont_mask_y, *, svm_model, mtf_seeds=None, recipe=None,
        keep_steps=False, mtf_structure=None,
    ):
        captured["seeds"] = mtf_seeds
        captured["structure"] = mtf_structure
        return _fake_result()

    monkeypatch.setattr(worker_module, "evaluate_slice_quality", fake_evaluate)
    window = MainWindowView()
    volume = _load_fake_volume(window)
    volume._slice_cache[0][:, 12] = 100.0  # a vertical "falx" at column 12
    _set_roi(window, "roi1", RoiRect(0, 0, 20, 20))
    _set_roi(window, "roi2", RoiRect(20, 20, 40, 40))
    window._confirm_rois_cb.setChecked(True)

    window._run_all_btn.click()
    window._scheduler.pool.waitForDone(5000)

    first_x, _first_y = captured["seeds"][0]
    assert first_x == 12
    assert captured["structure"] is None  # manual ROI on an unsegmented slice: no structure check


def test_roi_selection_and_run_all_are_logged_to_the_trace_panel(monkeypatch):
    monkeypatch.setattr(worker_module, "evaluate_slice_quality", lambda *a, **k: _fake_result())

    window = MainWindowView()
    window._show_trace_action.setChecked(True)
    _load_fake_volume(window)
    _set_roi(window, "roi1", RoiRect(0, 0, 10, 10))
    _set_roi(window, "roi2", RoiRect(20, 20, 40, 40))
    window._confirm_rois_cb.setChecked(True)

    window._run_all_btn.click()
    window._scheduler.pool.waitForDone(5000)
    for _ in range(20):
        QApplication.processEvents()

    trace_text = window._trace_log.toPlainText()
    assert "Falx ROI1 set" in trace_text
    assert "Falx ROI2 set" in trace_text
    assert "Run All: slice=0 (Falx)" in trace_text
    assert "Run All succeeded" in trace_text


def test_window_preset_defaults_to_brain_and_switching_updates_viewport():
    window = MainWindowView()
    _load_fake_volume(window)

    assert window._window_preset_combo.currentText() == "Brain"

    window._on_window_preset_clicked("Bone")

    assert window._roi_viewport._window_width == 3000.0
    assert window._roi_viewport._window_center == 400.0


def test_custom_window_action_disabled_before_a_study_is_loaded():
    window = MainWindowView()
    assert window._custom_window_action.isEnabled() is False


def test_custom_window_dialog_accept_applies_the_chosen_window(monkeypatch):
    from PySide6.QtWidgets import QDialog

    class _FakeDialog:
        def __init__(self, *args, **kwargs):
            pass

        def exec(self):
            return QDialog.DialogCode.Accepted

        def window(self):
            return (250.0, 60.0)

    monkeypatch.setattr(main_window_module, "CustomWindowDialog", _FakeDialog)

    window = MainWindowView()
    _load_fake_volume(window)

    window._on_custom_window_clicked()

    assert window._roi_viewport._window_width == 250.0
    assert window._roi_viewport._window_center == 60.0
    assert "WL: 60" in _overlay_text(window, "bottom_right")
    assert "WW: 250" in _overlay_text(window, "bottom_right")


def test_custom_window_dialog_cancel_leaves_the_window_unchanged(monkeypatch):
    from PySide6.QtWidgets import QDialog

    class _FakeDialog:
        def __init__(self, *args, **kwargs):
            pass

        def exec(self):
            return QDialog.DialogCode.Rejected

    monkeypatch.setattr(main_window_module, "CustomWindowDialog", _FakeDialog)

    window = MainWindowView()
    _load_fake_volume(window)
    original_width, original_center = window._roi_viewport.current_window()

    window._on_custom_window_clicked()

    assert window._roi_viewport.current_window() == (original_width, original_center)


def test_window_preset_dropdown_offers_all_standard_presets():
    window = MainWindowView()
    items = [window._window_preset_combo.itemText(i) for i in range(window._window_preset_combo.count())]
    assert items == list(WINDOW_PRESETS)
    assert "Brain" in items and "Subdural" in items and "Bone" in items and "Stroke" in items


def test_confirm_checkbox_disabled_until_slice_and_both_rois_present():
    window = MainWindowView()
    assert window._confirm_rois_cb.isEnabled() is False

    _load_fake_volume(window)
    _set_roi(window, "roi1", RoiRect(0, 0, 10, 10))
    assert window._confirm_rois_cb.isEnabled() is False  # only one ROI so far

    _set_roi(window, "roi2", RoiRect(20, 20, 30, 30))
    assert window._confirm_rois_cb.isEnabled() is True
    assert window._roi_position_items["falx.roi1"].text() == "0–10, 0–10"
    assert window._roi_size_items["falx.roi1"].text() == "10×10 px"
    assert window._roi_position_items["falx.roi2"].text() == "20–30, 20–30"
    assert window._roi_size_items["falx.roi2"].text() == "10×10 px"


def test_drawing_a_rectangle_shows_a_live_size_preview_without_finalizing_selection():
    window = MainWindowView()
    _load_fake_volume(window)

    window._roi_viewport.draft_changed.emit(RoiRect(10, 10, 25, 25))

    assert _overlay_text(window, "bottom_right").endswith("Selecting: 15×15 px")
    # Drawing must not count as a real ROI1/ROI2 selection yet - only
    # right-click > Save as does (`_save_draft_as`).
    assert "falx.roi1" not in window._roi_rects
    assert window._confirm_rois_cb.isEnabled() is False


def _menu_actions(menu):
    """{text: action} of a menu and its submenus (submenu entries as "Save as/<text>")."""
    out = {}
    for action in menu.actions():
        if action.menu():
            for sub in action.menu().actions():
                out[f"{action.text()}/{sub.text()}"] = sub
        elif action.text():
            out[action.text()] = action
    return out


def test_right_click_on_a_drawn_rectangle_offers_delete_and_save_as_the_four_roles():
    window = MainWindowView()
    _load_fake_volume(window)
    window._roi_viewport.set_draft(RoiRect(10, 10, 25, 25))

    actions = _menu_actions(window._roi_context_menu(""))

    assert {"Delete", "Save as/Falx", "Save as/White matter (next to falx)", "Save as/Septum",
            "Save as/Ventricle (next to septum)"} <= set(actions)
    actions["Save as/Ventricle (next to septum)"].trigger()
    assert window._roi_rects["septum.roi2"] == RoiRect(10, 10, 25, 25)
    assert not window._roi_viewport.has_draft()


def test_a_drawn_rectangle_can_be_deleted_with_the_button_the_key_or_the_menu():
    window = MainWindowView()
    _load_fake_volume(window)

    for delete in (
        lambda: window._delete_roi_btn.click(),
        lambda: window._roi_viewport.roi_delete_requested.emit(""),  # what the Del key sends
        lambda: _menu_actions(window._roi_context_menu(""))["Delete"].trigger(),
    ):
        window._roi_viewport.set_draft(RoiRect(10, 10, 25, 25))
        assert window._delete_roi_btn.isEnabled()
        delete()
        assert not window._roi_viewport.has_draft()
        assert not window._delete_roi_btn.isEnabled()


def test_right_click_save_as_moves_an_existing_roi_to_another_role():
    window = MainWindowView()
    _load_fake_volume(window)
    _set_roi(window, "falx.roi1", RoiRect(0, 0, 10, 10))

    actions = _menu_actions(window._roi_context_menu("falx.roi1"))
    assert not actions["Save as/Falx"].isEnabled()  # its own role
    actions["Save as/Septum"].trigger()

    assert window._roi_rects.get("septum.roi1") == RoiRect(0, 0, 10, 10)
    assert "falx.roi1" not in window._roi_rects


def test_confirmed_rois_have_a_disabled_right_click_menu():
    window = MainWindowView()
    _load_fake_volume(window)
    _set_roi(window, "roi1", RoiRect(0, 0, 10, 10))
    _set_roi(window, "roi2", RoiRect(20, 20, 30, 30))
    window._confirm_rois_cb.setChecked(True)

    actions = _menu_actions(window._roi_context_menu("falx.roi1"))

    assert not actions["Delete"].isEnabled()
    assert not actions["Save as/Septum"].isEnabled()
    assert any("confirmed" in text for text in actions)


def test_hovering_the_image_shows_cursor_position_and_value_in_the_overlay():
    window = MainWindowView()
    _load_fake_volume(window)

    window._roi_viewport.cursor_moved.emit((12, 34, -56.0))
    assert _overlay_text(window, "bottom_right").endswith("X: 12  Y: 34  Val: -56 HU")

    window._roi_viewport.cursor_moved.emit(None)
    assert "Val:" not in _overlay_text(window, "bottom_right")


def test_reselecting_an_roi_resets_confirmation():
    window = MainWindowView()
    _load_fake_volume(window)
    _set_roi(window, "roi1", RoiRect(0, 0, 10, 10))
    _set_roi(window, "roi2", RoiRect(20, 20, 30, 30))
    window._confirm_rois_cb.setChecked(True)

    _set_roi(window, "roi1", RoiRect(5, 5, 15, 15))

    assert window._confirm_rois_cb.isChecked() is False


def test_confirm_toggled_gates_run_all_button():
    window = MainWindowView()
    _load_fake_volume(window)
    _set_roi(window, "roi1", RoiRect(0, 0, 10, 10))
    _set_roi(window, "roi2", RoiRect(20, 20, 30, 30))

    assert window._run_all_btn.isEnabled() is False
    window._confirm_rois_cb.setChecked(True)
    assert window._run_all_btn.isEnabled() is True
    window._confirm_rois_cb.setChecked(False)
    assert window._run_all_btn.isEnabled() is False


def test_run_all_dispatches_worker_and_renders_real_result(monkeypatch):
    monkeypatch.setattr(worker_module, "evaluate_slice_quality", lambda *a, **k: _fake_result())

    window = MainWindowView()
    _load_fake_volume(window)
    _set_roi(window, "roi1", RoiRect(0, 0, 10, 10))
    _set_roi(window, "roi2", RoiRect(20, 20, 40, 40))
    window._confirm_rois_cb.setChecked(True)

    window._run_all_btn.click()
    window._scheduler.pool.waitForDone(5000)
    for _ in range(20):
        QApplication.processEvents()

    assert window._metric_status_items["mtf"].text() == "done"
    assert window._metric_value_items["mtf"].text() == "0.418 mm⁻¹"
    assert window._metric_value_items["contrast"].text() == "812.3 HU"
    assert window._result_verdict_label.text() == "Acceptable"
    assert window._result_decision_label.text() == "decision +1.42"
    assert window._feature_labels["mtf_area"].text() == "0.418"
    assert window._results_tab_btn.isEnabled() is True
    assert window._view_stack.currentIndex() == 1
    assert window._run_all_btn.isEnabled() is True

    # MTF plots the mean across edges (1 series), regardless of edge count; NPS has 1 curve.
    assert len(window._mtf_chart_view.chart().series()) == 1
    assert len(window._nps_chart_view.chart().series()) == 1

    mtf_points = window._mtf_chart_view.chart().series()[0].points()
    expected_mean = np.mean([[1.0, 1.0], [0.6, 0.5], [0.2, 0.1]], axis=1)
    assert [p.y() for p in mtf_points] == pytest.approx(list(expected_mean))

    # NPS must be ROI2 (homogeneous), PS must be ROI1 (structured) - these
    # were previously swapped (NPS chart showed ROI1 data).
    nps_points = window._nps_chart_view.chart().series()[0].points()
    assert [p.y() for p in nps_points] == pytest.approx([0.003, 0.0022, 0.0009])
    ps_points = window._ps_chart_view.chart().series()[0].points()
    assert [p.y() for p in ps_points] == pytest.approx([0.002, 0.0015, 0.0005])
    ratio_points = window._ratio_chart_view.chart().series()[0].points()
    assert [p.y() for p in ratio_points] == pytest.approx([2.0, 1.5, 1.0])

    # nps result label uses the true (ROI2/homogeneous) NPS, not ROI1's.
    assert window._metric_value_items["nps"].text() == "0.00291 HU²mm²"


def test_verdict_label_color_reflects_acceptable_vs_unacceptable(monkeypatch):
    window = MainWindowView()
    _load_fake_volume(window)
    _set_roi(window, "roi1", RoiRect(0, 0, 10, 10))
    _set_roi(window, "roi2", RoiRect(20, 20, 40, 40))
    window._confirm_rois_cb.setChecked(True)

    acceptable_result = _fake_result()
    monkeypatch.setattr(worker_module, "evaluate_slice_quality", lambda *a, **k: acceptable_result)
    window._run_all_btn.click()
    window._scheduler.pool.waitForDone(5000)
    for _ in range(20):
        QApplication.processEvents()
    assert window._result_verdict_label.text() == "Acceptable"
    assert "5fbf83" in window._result_verdict_label.styleSheet()  # GOOD

    unacceptable_result = _fake_result()
    unacceptable_result.predicted_label = 0
    monkeypatch.setattr(worker_module, "evaluate_slice_quality", lambda *a, **k: unacceptable_result)
    window._run_all_btn.click()
    window._scheduler.pool.waitForDone(5000)
    for _ in range(20):
        QApplication.processEvents()
    assert window._result_verdict_label.text() == "Unacceptable"
    assert "e0685c" in window._result_verdict_label.styleSheet()  # CRITICAL


def test_chart_axes_use_fixed_short_label_format():
    # QValueAxis's default (empty) label format lets Qt pick decimal
    # precision dynamically, which for real curve data can produce tick
    # labels too wide to fit - Qt then renders "..." instead. A fixed,
    # short format avoids that regardless of the data's actual range.
    window = MainWindowView()
    charts = (window._mtf_chart_view, window._nps_chart_view, window._ps_chart_view, window._ratio_chart_view)
    for chart_view in charts:
        for axis in chart_view.chart().axes():
            assert axis.labelFormat() != ""
            assert axis.tickCount() <= 6


def test_per_slice_starts_on_the_curves_and_the_switch_shows_which_view_is_active():
    # isVisible() always reports False for any un-shown window's children in
    # Qt, regardless of setVisible() calls - isHidden() reflects the actual
    # explicit visibility state we set, so that's what these tests check.
    window = MainWindowView()
    assert window._curves_view_btn.isChecked() and not window._steps_view_btn.isChecked()
    assert window._results_curves_stack.currentIndex() == 0 and window._feature_box.isHidden()

    window._steps_view_btn.click()
    assert window._steps_view_btn.isChecked() and not window._curves_view_btn.isChecked()
    assert window._results_curves_stack.currentWidget() is window._slice_steps_panel
    assert not window._feature_box.isHidden()
    assert MainWindowView()._curves_view_btn.isChecked()  # not remembered: a new window starts on the curves

    window._curves_view_btn.click()
    assert window._results_curves_stack.currentIndex() == 0 and window._feature_box.isHidden()


def test_editing_controls_and_run_all_hidden_on_results_page_metric_rows_stay_visible():
    window = MainWindowView()
    _load_fake_volume(window)

    # Viewer page: editing controls + Run All visible.
    window._show_view_page(0)
    assert window._viewer_controls_widget.isHidden() is False
    assert window._segmentation_roi_widget.isHidden() is False
    assert window._run_all_btn.isHidden() is False

    # Results page: no way to (re-)run metrics or change ROIs/windowing -
    # there's no ROI to run against from this page. The metrics table (a
    # sibling in the sidebar) is never touched here, so it stays visible
    # on both pages.
    window._results_tab_btn.setEnabled(True)
    window._show_view_page(1)
    assert window._viewer_controls_widget.isHidden() is True
    assert window._segmentation_roi_widget.isHidden() is True
    assert window._run_all_btn.isHidden() is True
    assert window._metrics_table.isHidden() is False


def test_run_all_surfaces_worker_error(monkeypatch):
    def _raise(*args, **kwargs):
        raise ValueError("ROI2 is smaller than the MTF active-contour mask")

    monkeypatch.setattr(worker_module, "evaluate_slice_quality", _raise)

    window = MainWindowView()
    _load_fake_volume(window)
    _set_roi(window, "roi1", RoiRect(0, 0, 10, 10))
    _set_roi(window, "roi2", RoiRect(20, 20, 40, 40))
    window._confirm_rois_cb.setChecked(True)

    window._run_all_btn.click()
    window._scheduler.pool.waitForDone(5000)
    for _ in range(20):
        QApplication.processEvents()

    assert window._metric_status_items["mtf"].text() == "error"
    assert "active-contour" in window._metric_status_items["mtf"].toolTip()
    assert window._run_all_btn.isEnabled() is True
    assert window._results_tab_btn.isEnabled() is False


def test_slice_slider_navigates_and_keeps_each_slices_rois_and_the_global_confirmation():
    # ROIs are per slice and one "Confirmed" covers every slice - navigating
    # shows the new slice's (here: no) ROIs without un-confirming, and
    # coming back shows the first slice's ROIs again.
    window = MainWindowView()
    _load_fake_volume(window, num_slices=5)

    _set_roi(window, "roi1", RoiRect(0, 0, 10, 10))
    _set_roi(window, "roi2", RoiRect(20, 20, 30, 30))
    window._confirm_rois_cb.setChecked(True)

    window._slice_slider.setValue(3)

    assert window.view_model.study.current_slice == 3
    assert window._slice_label.text() == "Slice 4 / 5"
    assert window._confirm_rois_cb.isChecked() is True
    assert window._roi_position_items["falx.roi1"].text() == "not set"
    assert window._roi_position_items["falx.roi2"].text() == "not set"

    window._slice_slider.setValue(0)

    assert window._roi_rects == {"falx.roi1": RoiRect(0, 0, 10, 10), "falx.roi2": RoiRect(20, 20, 30, 30)}


def test_confirming_locks_the_rois_and_unconfirming_unlocks_them():
    window = MainWindowView()
    _load_fake_volume(window)
    _set_roi(window, "roi1", RoiRect(0, 0, 10, 10))
    _set_roi(window, "roi2", RoiRect(20, 20, 30, 30))

    window._confirm_rois_cb.setChecked(True)
    assert window._roi_viewport._locked is True
    assert window._delete_roi_btn.isEnabled() is False

    window._confirm_rois_cb.setChecked(False)
    assert window._roi_viewport._locked is False


def _complete_a_run(window: MainWindowView, monkeypatch) -> None:
    """Drives a real Run All to completion (ROIs + worker), so `_last_result`/
    `_last_export_context` and the Results page are populated - the
    prerequisite state for Export tests.
    """
    monkeypatch.setattr(worker_module, "evaluate_slice_quality", lambda *a, **k: _fake_result())
    _set_roi(window, "roi1", RoiRect(0, 0, 10, 10))
    _set_roi(window, "roi2", RoiRect(20, 20, 40, 40))
    window._confirm_rois_cb.setChecked(True)
    window._run_all_btn.click()
    window._scheduler.pool.waitForDone(5000)
    for _ in range(20):
        QApplication.processEvents()


def test_close_study_resets_everything_to_the_no_study_state(monkeypatch):
    window = MainWindowView()
    _load_fake_volume(window, num_slices=3)
    _complete_a_run(window, monkeypatch)
    assert window._export_action.isEnabled() is True  # sanity: there's something to reset

    window._on_close_study_clicked()

    assert window.view_model.study.volume is None
    assert window._series_combo.count() == 0
    assert window._series_combo.isEnabled() is False
    assert window._roi_rects == {}
    assert window._roi_position_items["falx.roi1"].text() == "not set"
    assert window._confirm_rois_cb.isEnabled() is False
    assert window._delete_roi_btn.isEnabled() is False
    assert window._slice_slider.isEnabled() is False
    assert window._slice_label.text() == "No slices loaded"
    assert window._metric_status_items["mtf"].text() == "—"
    assert window._result_verdict_label.text() == "—"
    assert window._results_tab_btn.isEnabled() is False
    assert window._export_action.isEnabled() is False
    assert window.view_model.export.last_result is None
    assert window._close_study_action.isEnabled() is False
    assert window._reload_series_action.isEnabled() is False
    assert window._view_tags_action.isEnabled() is False
    assert window._view_stack.currentIndex() == 0


def test_study_and_view_menu_actions_enabled_only_after_a_volume_loads():
    window = MainWindowView()
    assert window._close_study_action.isEnabled() is False
    assert window._reload_series_action.isEnabled() is False
    assert window._view_tags_action.isEnabled() is False

    _load_fake_volume(window)

    assert window._close_study_action.isEnabled() is True
    assert window._reload_series_action.isEnabled() is True
    assert window._view_tags_action.isEnabled() is True


def test_reload_series_reloads_the_current_series(monkeypatch):
    monkeypatch.setattr(DicomVolume, "get_slice_hu", lambda self, index: np.zeros((8, 8), dtype=np.float64))
    window = MainWindowView()
    candidate = _fake_series_candidate("CT · 3 slice(s) · 8×8", num_slices=3)
    window._on_series_discovered([candidate])
    first_volume = window.view_model.study.volume

    window._on_reload_series_clicked()

    # A fresh DicomVolume, not the same object reused - a real reload.
    assert window.view_model.study.volume is not first_volume
    assert window.view_model.study.volume.num_slices == 3


def test_export_action_enabled_after_a_successful_run(monkeypatch):
    window = MainWindowView()
    _load_fake_volume(window)

    _complete_a_run(window, monkeypatch)

    assert window._export_action.isEnabled() is True
    # QAction.toolTip() falls back to the action's text() when the explicit
    # tooltip is empty (Qt behavior, not a bug) - the meaningful check is
    # that the disabled-state message is gone, not any particular string.
    assert window._export_action.toolTip() != "Run metrics first"
    assert window.view_model.export.last_result is not None
    assert window.view_model.export.last_context.slice_number == 1


def test_perform_export_writes_the_selected_files(monkeypatch, tmp_path):
    # `_perform_export` ends with a real (blocking, modal) QMessageBox on
    # success - patched to a no-op so this test doesn't hang waiting for a
    # click that will never come.
    monkeypatch.setattr(QMessageBox, "information", lambda *a, **k: None)
    window = MainWindowView()
    _load_fake_volume(window)
    _complete_a_run(window, monkeypatch)

    window._perform_export({"excel", "csv", "jpg"}, tmp_path / "Slice 1.xlsx")

    written = sorted(p.name for p in tmp_path.iterdir())
    assert written == [
        "Slice 1 - MTF.jpg", "Slice 1 - NPS.jpg", "Slice 1 - PS-MTF ratio.jpg", "Slice 1 - PS.jpg",
        "Slice 1.csv", "Slice 1.xlsx",
    ]


def test_perform_export_only_writes_the_selected_items(monkeypatch, tmp_path):
    monkeypatch.setattr(QMessageBox, "information", lambda *a, **k: None)
    window = MainWindowView()
    _load_fake_volume(window)
    _complete_a_run(window, monkeypatch)

    window._perform_export({"csv"}, tmp_path / "Slice 1.csv")

    assert [p.name for p in tmp_path.iterdir()] == ["Slice 1.csv"]


def test_roi_and_metrics_table_cells_are_center_aligned():
    window = MainWindowView()
    assert window._roi_position_items["falx.roi1"].textAlignment() == Qt.AlignmentFlag.AlignCenter
    assert window._roi_size_items["falx.roi1"].textAlignment() == Qt.AlignmentFlag.AlignCenter
    assert window._metric_value_items["mtf"].textAlignment() == Qt.AlignmentFlag.AlignCenter
    assert window._metric_status_items["mtf"].textAlignment() == Qt.AlignmentFlag.AlignCenter


def test_viewport_wheel_scroll_moves_slice_slider():
    window = MainWindowView()
    _load_fake_volume(window, num_slices=5)

    window._on_viewport_slice_scroll(1)
    assert window._slice_slider.value() == 1

    window._on_viewport_slice_scroll(1)
    assert window._slice_slider.value() == 2

    window._on_viewport_slice_scroll(-1)
    assert window._slice_slider.value() == 1

    # Clamps at the bounds instead of wrapping/erroring.
    for _ in range(10):
        window._on_viewport_slice_scroll(-1)
    assert window._slice_slider.value() == 0


# ---- Run All = batch QA over every slice with both ROIs ----


def _set_rois_on_slices(window: MainWindowView, slices: list[int]) -> None:
    for index in slices:
        window._slice_slider.setValue(index)
        _set_roi(window, "roi1", RoiRect(0, 0, 10, 10))
        _set_roi(window, "roi2", RoiRect(20, 20, 40, 40))


def _run_batch_to_completion(window: MainWindowView) -> None:
    window._confirm_rois_cb.setChecked(True)
    window._run_all_btn.click()
    for _ in range(200):
        window._scheduler.pool.waitForDone(100)
        QApplication.processEvents()
        if window.view_model.qa.jobs is None:
            break


def _result_with_label(label: int) -> SimpleNamespace:
    result = _fake_result()
    result.predicted_label = label
    return result


def test_run_all_computes_qa_for_every_slice_with_both_rois(monkeypatch):
    monkeypatch.setattr(worker_module, "evaluate_slice_quality", lambda *a, **k: _fake_result())
    window = MainWindowView()
    _load_fake_volume(window, num_slices=5)
    _set_rois_on_slices(window, [1, 3, 4])
    window._slice_slider.setValue(2)  # slice without ROIs is skipped

    _run_batch_to_completion(window)

    assert set(window.view_model.qa.results) == {(1, "falx"), (3, "falx"), (4, "falx")}
    assert window.view_model.results.slices == [1, 3, 4]
    assert window._view_stack.currentIndex() == 1
    assert window._qa_progress_label.text() == "QA done: 3 ROI pair(s)"
    assert window._study_table.rowCount() == 3


def test_results_page_scrolls_through_slices_and_follows_the_viewer(monkeypatch):
    monkeypatch.setattr(worker_module, "evaluate_slice_quality", lambda *a, **k: _fake_result())
    window = MainWindowView()
    _load_fake_volume(window, num_slices=5)
    _set_rois_on_slices(window, [1, 3])
    _run_batch_to_completion(window)

    window._results_slider.setValue(0)
    assert window.view_model.study.current_slice == 1
    assert window._results_slice_label.text().startswith("Slice 2 / 5")

    window._step_results_slice(1)
    assert window.view_model.study.current_slice == 3
    assert window.view_model.export.last_context.slice_number == 4  # export covers the shown slice


def test_study_verdict_is_the_majority_of_slice_verdicts(monkeypatch):
    labels = iter([1, 0, 1])
    monkeypatch.setattr(worker_module, "evaluate_slice_quality", lambda *a, **k: _result_with_label(next(labels)))
    window = MainWindowView()
    _load_fake_volume(window, num_slices=3)
    _set_rois_on_slices(window, [0, 1, 2])

    _run_batch_to_completion(window)

    assert window._study_verdict_label.text() == "Falx: Acceptable — 2 of 3 slice(s) acceptable"
    assert window._study_table.item(1, 1).text() == "Falx"
    assert window._study_table.item(1, 6).text() == "Unacceptable"


def test_a_failing_slice_does_not_stop_the_batch(monkeypatch):
    calls = iter([None, ValueError("no edge found"), None])

    def evaluate(*args, **kwargs):
        outcome = next(calls)
        if outcome is not None:
            raise outcome
        return _fake_result()

    monkeypatch.setattr(worker_module, "evaluate_slice_quality", evaluate)
    window = MainWindowView()
    _load_fake_volume(window, num_slices=3)
    _set_rois_on_slices(window, [0, 1, 2])

    _run_batch_to_completion(window)

    assert set(window.view_model.qa.results) == {(0, "falx"), (2, "falx")}
    assert window.view_model.qa.errors == {(1, "falx"): "no edge found"}
    assert window._qa_progress_label.text() == "QA done: 3 ROI pair(s), 1 failed"
    # Failed slices are left out of the results, only counted.
    assert window.view_model.results.slices == [0, 2]
    assert window._study_table.rowCount() == 2
    assert [window._study_table.item(row, 0).text() for row in range(2)] == ["1", "3"]
    assert window._study_verdict_label.text().endswith("(1 failed, not shown)")


def test_cancel_stops_the_batch_after_the_running_slice(monkeypatch):
    window = MainWindowView()

    def evaluate(*args, **kwargs):
        window.view_model.qa.cancel_requested = True  # as if Cancel was clicked while this slice ran
        return _fake_result()

    monkeypatch.setattr(worker_module, "evaluate_slice_quality", evaluate)
    _load_fake_volume(window, num_slices=3)
    _set_rois_on_slices(window, [0, 1, 2])

    _run_batch_to_completion(window)

    assert set(window.view_model.qa.results) == {(0, "falx")}
    assert window._qa_progress_label.text() == "QA cancelled: 1/3 ROI pair(s)"
    assert window._cancel_run_btn.isHidden()


def test_run_metrics_box_shows_study_averages_on_the_summary_tab_and_slice_values_per_slice(monkeypatch):
    areas = iter([0.2, 0.4])

    def evaluate(*args, **kwargs):
        result = _fake_result()
        result.mtf_area = next(areas)
        return result

    monkeypatch.setattr(worker_module, "evaluate_slice_quality", evaluate)
    window = MainWindowView()
    _load_fake_volume(window, num_slices=2)
    _set_rois_on_slices(window, [0, 1])
    _run_batch_to_completion(window)
    window._results_slider.setValue(0)  # per-slice tab on slice 1

    assert window._metric_value_items["mtf"].text() == "0.200 mm⁻¹"

    window._results_tabs.setCurrentIndex(1)  # Study summary
    assert window._metric_value_items["mtf"].text() == "0.300 mm⁻¹"
    assert window._metric_status_items["mtf"].text() == "mean"
    assert "2 slice(s)" in window._metric_status_items["mtf"].toolTip()

    window._results_tabs.setCurrentIndex(0)  # back to Per slice
    assert window._metric_value_items["mtf"].text() == "0.200 mm⁻¹"

    window._results_tabs.setCurrentIndex(1)
    window._show_view_page(0)  # the image viewer always shows its slice's values
    assert window._metric_value_items["mtf"].text() == "0.200 mm⁻¹"


def test_study_export_writes_the_summary_and_per_slice_files(monkeypatch, tmp_path):
    import openpyxl

    monkeypatch.setattr(worker_module, "evaluate_slice_quality", lambda *a, **k: _fake_result())
    monkeypatch.setattr(QMessageBox, "information", lambda *a, **k: None)
    window = MainWindowView()
    _load_fake_volume(window, num_slices=3)
    _set_rois_on_slices(window, [0, 2])
    _run_batch_to_completion(window)

    window._perform_study_export({"excel", "csv"}, tmp_path / "Study")

    written = sorted(path.name for path in tmp_path.iterdir())
    assert written == ["Study - per slice.csv", "Study - summary.csv", "Study.xlsx"]
    workbook = openpyxl.load_workbook(tmp_path / "Study.xlsx")
    assert workbook.sheetnames == ["Summary", "Per slice"]
    assert len(list(workbook["Per slice"].iter_rows())) == 3  # header + slices 1 and 3


def test_the_export_is_saved_under_the_chosen_name_and_the_folder_is_remembered(monkeypatch, tmp_path):
    from pathlib import Path

    from PySide6.QtWidgets import QFileDialog

    from bictool.infrastructure.ui_settings import EXPORT_FOLDER

    window = MainWindowView()
    _load_fake_volume(window)
    _complete_a_run(window, monkeypatch)
    asked = []

    def save_dialog(_parent, _title, suggested, file_filter):
        asked.append((Path(suggested).name, file_filter))
        return str(tmp_path / "chosen.xlsx"), file_filter

    monkeypatch.setattr(QFileDialog, "getSaveFileName", save_dialog)

    path = window._ask_export_path("slice", {"excel", "csv"})

    assert path == tmp_path / "chosen.xlsx"
    assert asked[0][0].startswith("BICTool slice 1 ") and asked[0][0].endswith(".xlsx")
    assert window._ui_settings.get_str(EXPORT_FOLDER) == str(tmp_path)
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a: ("", ""))
    assert window._ask_export_path("slice", {"excel"}) is None  # cancelled


def _menu_path(window, *titles):
    """The QMenu reached by following menu titles from the menu bar."""
    menu = None
    actions = window.menuBar().actions()
    for title in titles:
        action = next(a for a in actions if a.text().replace("&", "") == title)
        menu = action.menu()
        actions = menu.actions()
    return menu


def test_qa_recipe_defaults_to_physical_and_is_tucked_under_tools_advanced():
    window = MainWindowView()

    recipe_menu = _menu_path(window, "Tools", "Advanced", "QA Recipe")
    assert [a.text() for a in recipe_menu.actions()] == ["Physical (default)", "Octave (legacy)"]
    assert window.view_model.qa.recipe == "physical"
    assert window._qa_recipe_actions["physical"].isChecked() is True
    assert window._qa_recipe_label.isHidden() is True


def test_legacy_recipe_is_passed_to_the_worker_and_shown_in_the_metrics_box(monkeypatch):
    received = []
    monkeypatch.setattr(worker_module, "evaluate_slice_quality", lambda *a, **k: received.append(k) or _fake_result())
    window = MainWindowView()
    _load_fake_volume(window, num_slices=1)
    _set_rois_on_slices(window, [0])

    window._qa_recipe_actions["octave"].trigger()
    _run_batch_to_completion(window)

    assert received[0]["recipe"] == "octave"
    assert window._qa_recipe_label.isHidden() is False


def test_switching_recipe_clears_results_of_the_other_recipe(monkeypatch):
    monkeypatch.setattr(worker_module, "evaluate_slice_quality", lambda *a, **k: _fake_result())
    window = MainWindowView()
    _load_fake_volume(window, num_slices=2)
    _set_rois_on_slices(window, [0, 1])
    _run_batch_to_completion(window)
    assert window.view_model.qa.results

    window._qa_recipe_actions["octave"].trigger()

    assert window.view_model.qa.results == {}
    assert window._confirm_rois_cb.isChecked() is True  # ROIs stay confirmed, Run All recomputes


def test_recipe_cannot_be_switched_while_a_batch_runs(monkeypatch):
    window = MainWindowView()

    def evaluate(*args, **kwargs):
        assert window._qa_recipe_actions["octave"].isEnabled() is False
        return _fake_result()

    monkeypatch.setattr(worker_module, "evaluate_slice_quality", evaluate)
    _load_fake_volume(window, num_slices=1)
    _set_rois_on_slices(window, [0])
    _run_batch_to_completion(window)

    assert window._qa_recipe_actions["octave"].isEnabled() is True


# ---- GUI clean-up 2026-10-07: settings, hover help, numbered steps ----

def test_trace_log_is_off_by_default_and_the_choice_is_remembered():
    import logging

    window = MainWindowView()
    assert window._trace_box.isHidden()
    assert window._log_handler not in logging.getLogger("bictool").handlers  # nothing is fed to a hidden panel

    window._show_trace_action.setChecked(True)
    assert not window._trace_box.isHidden()
    assert window._log_handler in logging.getLogger("bictool").handlers

    again = MainWindowView()  # a new session
    assert again._show_trace_action.isChecked() and not again._trace_box.isHidden()


def test_roi_details_are_hidden_by_default_and_remembered():
    window = MainWindowView()
    assert window._roi_table.isHidden()

    window._roi_details_btn.click()
    assert not window._roi_table.isHidden()
    assert window._roi_details_btn.text() == "▾ ROI details"

    assert not MainWindowView()._roi_table.isHidden()


def test_controls_show_help_and_why_they_are_disabled():
    from bictool.presentation.viewmodels.help_text import TOOLTIPS

    window = MainWindowView()

    tip = window._run_all_btn.toolTip()
    assert tip.startswith(TOOLTIPS["_run_all_btn"]) and tip.endswith("Confirm ROIs to enable")
    assert window._export_action.toolTip().startswith(TOOLTIPS["_export_action"])


def test_sidebar_steps_are_numbered_and_the_results_page_shows_slice_results():
    from PySide6.QtWidgets import QGroupBox

    from bictool.presentation.views.collapsible_section import CollapsibleSection

    window = MainWindowView()
    titles = [box.title() for box in window.findChildren(QGroupBox) + window.findChildren(CollapsibleSection)]
    assert {"1. Segmentation", "Anatomical", "WM/GM separation", "2. ROI Selection", "Confirm ROIs",
            "3. Run QA metrics"} <= set(titles)

    window._show_view_page(1)
    assert window._metrics_box.title() == "Slice results"
    assert window._series_box.isHidden()
    window._results_tabs.setCurrentIndex(1)
    assert window._metrics_box.title() == "Study averages (Falx)"
    window._show_view_page(0)
    assert window._metrics_box.title() == "3. Run QA metrics" and not window._series_box.isHidden()


def test_auto_roi_button_reads_start_or_remove():
    window = MainWindowView()
    assert window._auto_rois_btn.text() == "Start auto ROIs"

    window._auto_rois_btn.setEnabled(True)
    window._auto_rois_btn.setChecked(True)
    assert window._auto_rois_btn.text() == "Remove auto ROIs"
    window._auto_rois_btn.setChecked(False)
    assert window._auto_rois_btn.text() == "Start auto ROIs"


def test_sidebar_steps_collapse_and_stay_collapsed():
    window = MainWindowView()
    assert window._roi_section.is_expanded()

    window._roi_section.header().click()
    assert not window._roi_section.is_expanded()
    assert window._auto_rois_btn.isHidden() is False and not window._auto_rois_btn.isVisibleTo(window._roi_section)

    assert not MainWindowView()._roi_section.is_expanded()  # remembered


def test_viewer_and_sidebar_are_split_by_a_draggable_handle_that_is_remembered():
    window = MainWindowView()
    window.resize(1600, 1000)
    window.show()
    window._main_splitter.setSizes([900, 700])
    window._main_splitter.splitterMoved.emit(900, 1)

    again = MainWindowView()
    again.resize(1600, 1000)
    again.show()
    assert again._main_splitter.sizes()[1] > 600


def test_boxes_and_tabs_describe_themselves_on_hover():
    from bictool.presentation.viewmodels.help_text import TAB_TOOLTIPS, TOOLTIPS

    window = MainWindowView()

    assert window._wmgm_box.toolTip() == TOOLTIPS["_wmgm_box"]
    assert window._roi_section.toolTip() == TOOLTIPS["_roi_section"]
    assert window._roi_section.header().toolTip().startswith(TOOLTIPS["_roi_section"])
    assert window._results_tabs.tabToolTip(1) == TAB_TOOLTIPS["summary"]


def test_progress_and_status_live_in_the_status_bar():
    window = MainWindowView()
    status_bar = window.statusBar()

    for item in (window._study_activity, window._segmentation_activity, window._qa_activity, window._roi_activity):
        assert item.parentWidget() is not None and status_bar.isAncestorOf(item)
    assert status_bar.isAncestorOf(window._cancel_segment_volume_btn)
    assert status_bar.isAncestorOf(window._cancel_run_btn)
    assert not window._metrics_box.isAncestorOf(window._qa_progress_label)


def test_sidebar_and_its_boxes_grow_with_the_window_and_the_handle():
    from PySide6.QtWidgets import QApplication

    window = MainWindowView()
    window.resize(1400, 900)
    window.show()
    QApplication.processEvents()
    narrow = window._roi_section.width()

    window._main_splitter.setSizes([700, 700])  # handle dragged left
    QApplication.processEvents()
    assert window._roi_section.width() > narrow + 200

    before = window._main_splitter.sizes()[1]
    window.resize(2400, 1300)  # full screen
    QApplication.processEvents()
    assert window._main_splitter.sizes()[1] > before


def test_the_wm_gm_box_tells_the_user_to_choose_the_method():
    window = MainWindowView()
    hint = window._brain_tissue_method_hint.text()
    assert "Choose the method that fits this scan best" in hint
