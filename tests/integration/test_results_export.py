import csv
from types import SimpleNamespace

import openpyxl
import pytest

from bictool.integration.results_export import (
    ExportContext,
    StudyExportContext,
    export_metrics_csv,
    export_metrics_excel,
    export_study_csv,
    export_study_excel,
)


def _fake_result() -> SimpleNamespace:
    return SimpleNamespace(
        mtf_area=0.418,
        nps_roi2=SimpleNamespace(var_nps=0.00291),
        ratio=SimpleNamespace(ratio_lf_hf=1.86),
        contrast_roi1=812.3,
        decision_value=1.42,
        predicted_label=1,
        feature_vector=[0.418, 0.00142, 0.00187, 12.4, 1.86, -0.031, 0.34],
    )


def _fake_context() -> ExportContext:
    return ExportContext(
        series_label="CT · 467 slice(s) · 512×512",
        slice_number=12,
        slice_count=467,
        pixel_spacing_mm=0.4880,
        acquisition_date="20260101",
    )


def test_export_metrics_csv_writes_the_expected_rows(tmp_path):
    path = tmp_path / "result.csv"

    export_metrics_csv(_fake_result(), _fake_context(), path)

    with open(path, newline="", encoding="ascii") as f:
        rows = [tuple(row) for row in csv.reader(f)]

    assert ("Slice", "12/467") in rows
    assert ("MTF area (mm^-1)", "0.4180") in rows
    assert ("Verdict", "Acceptable") in rows
    assert ("mtf_area", "0.418") in rows


def test_export_metrics_csv_never_contains_patient_identifiers(tmp_path):
    # NFR-06/GDPR - exports must stay de-identified even though the live GUI
    # shows patient name/ID.
    path = tmp_path / "result.csv"

    export_metrics_csv(_fake_result(), _fake_context(), path)

    text = path.read_text(encoding="ascii")
    assert "patient" not in text.lower()


def test_export_metrics_csv_marks_unacceptable_verdict(tmp_path):
    path = tmp_path / "result.csv"
    result = _fake_result()
    result.predicted_label = 0

    export_metrics_csv(result, _fake_context(), path)

    with open(path, newline="", encoding="ascii") as f:
        rows = [tuple(row) for row in csv.reader(f)]
    assert ("Verdict", "Unacceptable") in rows


def test_export_metrics_excel_writes_the_expected_rows(tmp_path):
    path = tmp_path / "result.xlsx"

    export_metrics_excel(_fake_result(), _fake_context(), path)

    workbook = openpyxl.load_workbook(path)
    sheet = workbook.active
    rows = [tuple(cell.value for cell in row) for row in sheet.iter_rows()]

    assert ("Slice", "12/467") in rows
    assert ("Verdict", "Acceptable") in rows
    assert any(r[0] == "MTF area (mm^-1)" for r in rows)


def _study_results() -> dict:
    """Results per structure (ROI pair), keyed by 0-based slice index."""
    first = _fake_result()
    second = _fake_result()
    second.mtf_area = 0.218
    second.predicted_label = 0
    third = _fake_result()
    septum = _fake_result()
    septum.mtf_area = 0.9
    return {"Falx": {11: first, 12: second, 20: third}, "Septum": {12: septum}}


def _study_context() -> StudyExportContext:
    return StudyExportContext(
        series_label="CT head",
        slice_count=467,
        pixel_spacing_mm=0.4880,
        acquisition_date="20260101",
        failed_slices=["14 (Falx)", "30 (Septum)"],
    )


def _read_csv(path) -> list[list[str]]:
    with open(path, newline="", encoding="ascii") as f:
        return list(csv.reader(f))


def test_export_study_csv_writes_the_averages_and_a_row_per_slice(tmp_path):
    summary_path, slices_path = tmp_path / "summary.csv", tmp_path / "slices.csv"

    export_study_csv(_study_results(), _study_context(), summary_path, slices_path)

    summary = _read_csv(summary_path)
    assert ["Slices analysed", "3"] in summary  # slices 12, 13, 21 (13 has both pairs)
    assert ["Failed slices", "14 (Falx), 30 (Septum)"] in summary
    # One block per structure - never averaged together.
    falx = summary.index(["Falx"])
    septum = summary.index(["Septum"])
    assert summary[falx + 1] == ["Study verdict", "Acceptable"]
    assert summary[falx + 2] == ["Acceptable slices", "2 of 3"]
    falx_mtf = next(row for row in summary[falx:septum] if row and row[0] == "MTF area (mm^-1)")
    assert float(falx_mtf[1]) == pytest.approx((0.418 + 0.218 + 0.418) / 3)
    assert falx_mtf[3] == "3"
    septum_mtf = next(row for row in summary[septum:] if row and row[0] == "MTF area (mm^-1)")
    assert float(septum_mtf[1]) == pytest.approx(0.9)

    slices = _read_csv(slices_path)
    assert slices[0][:3] == ["Slice", "Structure", "MTF area (mm^-1)"]
    assert [row[:2] for row in slices[1:]] == [["12", "Falx"], ["13", "Falx"], ["13", "Septum"], ["21", "Falx"]]
    assert slices[2][7] == "Unacceptable"


def test_export_study_excel_has_a_summary_and_a_per_slice_sheet(tmp_path):
    path = tmp_path / "study.xlsx"

    export_study_excel(_study_results(), _study_context(), path)

    workbook = openpyxl.load_workbook(path)
    assert workbook.sheetnames == ["Study summary", "Per slice"]
    summary = list(workbook["Study summary"].iter_rows(values_only=True))
    assert ("Study verdict", "Acceptable") in [row[:2] for row in summary]
    assert any(row[0] == "QA recipe" for row in summary)
    per_slice = list(workbook["Per slice"].iter_rows(values_only=True))
    assert len(per_slice) == 5  # header + 3 falx slices + 1 septum slice


def test_export_study_never_contains_patient_identifiers(tmp_path):
    summary_path, slices_path = tmp_path / "summary.csv", tmp_path / "slices.csv"

    export_study_csv(_study_results(), _study_context(), summary_path, slices_path)

    text = summary_path.read_text(encoding="ascii") + slices_path.read_text(encoding="ascii")
    assert "patient" not in text.lower()


def _dataset_volumes():
    from bictool.domain.results.volume_qa import StructureQa, VolumeQa
    from bictool.domain.roi_selection.structured_roi import RoiCandidate

    def structure_qa(name, labels):
        qa = StructureQa(name)
        for index, label in enumerate(labels):
            result = _fake_result()
            result.predicted_label = label
            result.mtf_seed = (6, 7)
            qa.results[index + 10] = result
            qa.rois[index + 10] = (RoiCandidate(100, 200, 26, name), RoiCandidate(100, 230, 26, "White Matter"))
        qa.errors[30] = "No edge columns survived bump removal"
        return qa

    good = VolumeQa(
        group="CanonPrime",
        name="20250911_144757 - bad",
        series_label="Head 4mm (40 slices, 4.0 mm)",
        slice_count=40,
        segmented_slices=38,
        pixel_spacing_mm=0.45,
        structures={
            "Septum Pellucidum": structure_qa("Septum Pellucidum", [1, 0, 0]),
            "Dural Folds": structure_qa("Dural Folds", [1, 1]),
        },
    )
    broken = VolumeQa(group="Siemens ER", name="broken study", error="ValueError: no series")
    return [good, broken]


def test_export_dataset_excel_has_a_summary_row_per_volume_and_structure_and_a_sheet_per_volume(tmp_path):
    from bictool.integration.results_export import export_dataset_excel

    path = tmp_path / "dataset.xlsx"
    export_dataset_excel(_dataset_volumes(), path)

    workbook = openpyxl.load_workbook(path)
    assert workbook.sheetnames == ["Summary", "1 CanonPrime bad", "2 Siemens ER broken study"]
    rows = list(workbook["Summary"].iter_rows(values_only=True))
    assert "untrusted" in rows[1][0]
    assert rows[2][0].startswith("QA recipe: octave")  # the fixtures carry no recipe -> Octave-era
    header_index = next(i for i, row in enumerate(rows) if row[0] == "Dataset")
    header = rows[header_index]
    body = rows[header_index + 1 :]
    structure_column = header.index("Structure")
    assert [row[structure_column] for row in body] == ["Septum Pellucidum", "Dural Folds", None]
    verdict_column = header.index("Verdict (legacy chest SVM, untrusted)")
    assert body[0][verdict_column] == "Unacceptable"  # septum: 1 of 3 - never mixed with dural folds
    assert body[1][verdict_column] == "Acceptable"
    assert body[0][header.index("Failed slices")] == "31"
    assert body[2][header.index("Note")] == "ValueError: no series"


def test_export_dataset_csv_writes_the_summary_and_one_long_per_slice_table(tmp_path):
    from bictool.integration.results_export import export_dataset_csv

    summary_path, slices_path = tmp_path / "summary.csv", tmp_path / "slices.csv"
    export_dataset_csv(_dataset_volumes(), summary_path, slices_path)

    summary = _read_csv(summary_path)
    assert summary[0][:3] == ["Dataset", "Volume", "Structure"]
    assert len(summary) == 1 + 3
    slices = _read_csv(slices_path)
    assert slices[0][:4] == ["Dataset", "Volume", "Structure", "Slice"]
    assert len(slices) == 1 + 5  # 3 septum + 2 dural fold slices, failed ones not listed
    assert slices[1][:4] == ["CanonPrime", "20250911_144757 - bad", "Septum Pellucidum", "11"]


# --- Physical recipe: study-level head-CT verdict ------------------------------

def _physical_results(n=8):
    """Falx results with features of a typical good study (+ one noisy slice) and one septum result."""
    import numpy as np

    good = np.array([0.51, 20.0, 100.0, 70.0, 11.0, -1.0, 0.1])

    def result(features):
        r = _fake_result()
        r.recipe = "physical"
        r.feature_vector = np.array(features, dtype=float)
        r.nps_f_av, r.nps_f_peak = 0.2, 0.1
        return r

    falx = {i: result(good * (1 + 0.01 * (i % 3))) for i in range(n)}
    if 3 in falx:
        falx[3].feature_vector[1] = 80.0
    return {"Falx": falx, "Septum": {4: result(good)}}


def _physical_context(results):
    from bictool.domain.results.head_ct_verdict import evaluate_study_results
    from bictool.integration.model_files import default_head_ct_model

    verdict = evaluate_study_results(default_head_ct_model(), results["Falx"], results["Septum"])
    return verdict, StudyExportContext(
        series_label="CT", slice_count=40, pixel_spacing_mm=0.5, acquisition_date=None,
        failed_slices=[], head_ct_verdict=verdict,
    )


def test_physical_study_summary_has_the_verdict_with_intervals_and_the_provisional_note(tmp_path):
    results = _physical_results()
    verdict, context = _physical_context(results)

    export_study_csv(results, context, tmp_path / "s.csv", tmp_path / "p.csv")
    summary = _read_csv(tmp_path / "s.csv")

    assert ["Verdict", verdict.verdict] in summary
    ci = next(row for row in summary if row[:1] == ["95 % CI"])
    assert (float(ci[1]), float(ci[2])) == pytest.approx((verdict.probability_low, verdict.probability_high), abs=1e-3)
    assert ["Falx slices used", "8"] in summary
    assert ["Model", "head-ct-1"] in summary
    assert any(row[:1] == ["Feature (falx median)"] for row in summary)
    assert any(row[:1] == ["Artefact indicator (provisional)"] for row in summary)
    assert any(row[:1] == ["Slices with noise > 3x median"] and row[1] == "4" for row in summary)  # 1-based
    assert any(row[:1] == ["Note"] and "Provisional model" in row[1] for row in summary)
    # The legacy majority vote is not reported for the physical recipe.
    assert not any(row[:1] == ["Study verdict"] for row in summary)
    assert not any(row[:1] == ["Decision value"] for row in summary)


def test_physical_slice_table_marks_the_legacy_svm_and_adds_the_slice_score(tmp_path):
    results = _physical_results()
    verdict, context = _physical_context(results)

    export_study_csv(results, context, tmp_path / "s.csv", tmp_path / "p.csv")
    rows = _read_csv(tmp_path / "p.csv")
    header = rows[0]

    assert "Verdict (legacy chest SVM, untrusted)" in header
    assert header[-4:] == ["Slice score (indicative)", "Noise outlier", "Possible artefact", "Basal ganglia (provisional)"]
    by_slice = {(row[0], row[1]): row for row in rows[1:]}
    noisy = by_slice[("4", "Falx")]
    assert noisy[-3] == "yes" and 0.0 <= float(noisy[-4]) <= 1.0
    assert by_slice[("5", "Septum")][-4:-2] == ["", ""]  # no slice score for the septum


def test_physical_study_without_a_verdict_says_why(tmp_path):
    results = _physical_results(n=3)
    _, context = _physical_context(results)

    export_study_csv(results, context, tmp_path / "s.csv", tmp_path / "p.csv")
    summary = _read_csv(tmp_path / "s.csv")

    verdict_row = next(row for row in summary if row[:1] == ["Verdict"])
    assert "not available" in verdict_row[1] and "(this study: 3)" in verdict_row[1]


def test_legacy_recipe_study_summary_is_unchanged(tmp_path):
    results = {"Falx": {0: _fake_result()}}
    context = StudyExportContext("CT", 40, 0.5, None, [])

    export_study_csv(results, context, tmp_path / "s.csv", tmp_path / "p.csv")

    assert ["Study verdict", "Acceptable"] in _read_csv(tmp_path / "s.csv")
    assert _read_csv(tmp_path / "p.csv")[0][-3:] == ["loc_max", "Possible artefact", "Basal ganglia (provisional)"]


def test_physical_study_excel_contains_the_verdict(tmp_path):
    results = _physical_results()
    verdict, context = _physical_context(results)

    export_study_excel(results, context, tmp_path / "s.xlsx")
    sheet = openpyxl.load_workbook(tmp_path / "s.xlsx")["Study summary"]

    assert ["Verdict", verdict.verdict] in [[c for c in row[:2]] for row in sheet.iter_rows(values_only=True)]


def test_physical_single_slice_export_marks_the_chest_svm_verdict_as_untrusted(tmp_path):
    result = _physical_results(n=1)["Falx"][0]

    export_metrics_csv(result, _fake_context(), tmp_path / "r.csv")
    rows = _read_csv(tmp_path / "r.csv")

    assert ["Verdict (legacy chest SVM, untrusted)", "Acceptable"] in rows
    assert not any(row[:1] == ["Verdict"] for row in rows)


def test_dataset_summary_notes_slices_skipped_for_too_little_brain(tmp_path):
    from bictool.domain.results.volume_qa import StructureQa, VolumeQa
    from bictool.integration.results_export import export_dataset_csv

    qa = StructureQa("Dural Folds", results={}, errors={}, skipped=[3, 35])
    volume = VolumeQa("Scanner", "study", slice_count=40, segmented_slices=40, structures={"Dural Folds": qa})

    export_dataset_csv([volume], tmp_path / "s.csv", tmp_path / "p.csv")
    rows = _read_csv(tmp_path / "s.csv")

    note = rows[1][rows[0].index("Note")]
    assert "Skipped (< 20% of largest brain area): 4, 36" in note


def _physical_dataset_volume():
    from bictool.domain.results.head_ct_verdict import evaluate_study_results
    from bictool.domain.results.volume_qa import StructureQa, VolumeQa
    from bictool.domain.roi_selection.structured_roi import RoiCandidate
    from bictool.integration.model_files import default_head_ct_model

    results = _physical_results()
    structures = {}
    for structure, key in (("Dural Folds", "Falx"), ("Septum Pellucidum", "Septum")):
        qa = StructureQa(structure)
        for index, result in results[key].items():
            result.mtf_seed = (6, 7)
            qa.results[index] = result
            qa.rois[index] = (RoiCandidate(100, 200, 26, structure), RoiCandidate(100, 230, 26, "White Matter"))
        structures[structure] = qa
    verdict = evaluate_study_results(default_head_ct_model(), results["Falx"], results["Septum"])
    volume = VolumeQa("PCCT", "study", slice_count=40, segmented_slices=40, structures=structures, head_ct_verdict=verdict)
    return volume, verdict


def test_dataset_summary_carries_the_head_ct_verdict_on_the_falx_row(tmp_path):
    from bictool.integration.results_export import export_dataset_csv

    volume, verdict = _physical_dataset_volume()
    legacy = _dataset_volumes()[0]  # octave-era fixture

    export_dataset_csv([volume, legacy], tmp_path / "s.csv", tmp_path / "p.csv")
    rows = _read_csv(tmp_path / "s.csv")
    header = rows[0]
    by_row = {(row[1], row[2]): row for row in rows[1:]}
    falx, septum = by_row[("study", "Dural Folds")], by_row[("study", "Septum Pellucidum")]

    col = header.index("Study verdict (head-CT model)")
    assert falx[col] == verdict.verdict
    assert float(falx[header.index("P 95 % CI low")]) == pytest.approx(verdict.probability_low, abs=1e-4)
    assert falx[header.index("Falx slices used")] == "8"
    assert falx[header.index("Noise outlier slices")] == "4"
    assert falx[header.index("Model")] == "head-ct-1"
    assert septum[col] == ""  # the verdict belongs to the falx
    assert by_row[("20250911_144757 - bad", "Dural Folds")][col] == "not available (legacy recipe)"
    assert "Verdict (legacy chest SVM, untrusted)" in header


def test_dataset_slice_table_has_the_indicative_score_for_falx_slices(tmp_path):
    from bictool.integration.results_export import export_dataset_csv

    volume, verdict = _physical_dataset_volume()

    export_dataset_csv([volume], tmp_path / "s.csv", tmp_path / "p.csv")
    rows = _read_csv(tmp_path / "p.csv")
    header = rows[0]
    score, outlier = header.index("Slice score (indicative)"), header.index("Noise outlier")
    falx = {row[3]: row for row in rows[1:] if row[2] == "Dural Folds"}
    septum = [row for row in rows[1:] if row[2] == "Septum Pellucidum"]

    assert 0.0 <= float(falx["4"][score]) <= 1.0 and falx["4"][outlier] == "yes"
    assert falx["1"][outlier] == ""
    assert septum[0][score] == "" and septum[0][outlier] == ""


def test_dataset_volume_sheet_has_the_verdict_block(tmp_path):
    from bictool.integration.results_export import export_dataset_excel

    volume, verdict = _physical_dataset_volume()

    export_dataset_excel([volume], tmp_path / "d.xlsx")
    sheet = openpyxl.load_workbook(tmp_path / "d.xlsx").worksheets[1]
    rows = [list(row[:2]) for row in sheet.iter_rows(values_only=True)]

    assert ["Study verdict (head-CT model, falx)", None] in rows
    assert ["Verdict", verdict.verdict] in rows


def test_study_and_dataset_exports_list_the_possible_artefacts(tmp_path):
    from bictool.domain.artefacts.artefact_flags import ArtefactFlags
    from bictool.domain.results.volume_qa import StructureQa, VolumeQa
    from bictool.integration.results_export import export_dataset_csv

    flags = ArtefactFlags(metal_slices=(13, 14), metal_max_hu=3071.0, band_depths={4: 5.2, 5: 1.0})
    context = StudyExportContext("CT", 40, 0.5, None, [], artefacts=flags)
    export_study_csv(_study_results(), context, tmp_path / "s.csv", tmp_path / "p.csv")
    summary = _read_csv(tmp_path / "s.csv")
    assert ["Metal (> 2800 HU) on slices", "14, 15", "highest 3071 HU"] in summary
    assert ["Brainstem dark band >= 3.5 HU (provisional) on slices", "5"] in summary
    assert ["5", "5.20"] in summary and ["6", "1.00"] in summary

    structures = {name: StructureQa(name) for name in ("Dural Folds", "Septum Pellucidum")}
    volume = VolumeQa("Scanner", "study", slice_count=40, segmented_slices=40, structures=structures, artefact_flags=flags)
    export_dataset_csv([volume], tmp_path / "ds.csv", tmp_path / "dp.csv")
    rows = _read_csv(tmp_path / "ds.csv")
    metal, band = rows[0].index("Metal slices"), rows[0].index("Brainstem dark band slices (provisional)")
    assert (rows[1][metal], rows[1][band]) == ("14, 15", "5")
    assert (rows[2][metal], rows[2][band]) == ("", "")  # once per study


def test_the_exported_slice_table_names_artefacts_on_each_slice(tmp_path):
    from bictool.domain.artefacts.artefact_flags import ArtefactFlags

    flags = ArtefactFlags(metal_slices=(12,), metal_max_hu=3071.0, band_depths={12: 4.0})
    context = StudyExportContext("CT", 40, 0.5, None, [], artefacts=flags)
    export_study_csv(_study_results(), context, tmp_path / "s.csv", tmp_path / "p.csv")
    rows = _read_csv(tmp_path / "p.csv")
    column = rows[0].index("Possible artefact")
    by_slice = {(row[0], row[1]): row[column] for row in rows[1:]}
    assert by_slice[("13", "Falx")] == "metal; possible brainstem band"
    assert by_slice[("12", "Falx")] == ""


def test_exports_list_the_basal_ganglia_visibility(tmp_path):
    from bictool.domain.metrics.basal_ganglia import BasalGangliaSlice
    from bictool.domain.results.volume_qa import StructureQa, VolumeQa
    from bictool.integration.results_export import export_dataset_csv

    slices = {
        12: BasalGangliaSlice(0.9, {"Lentiform Nucleus": 6.5, "Caudate Nucleus": 7.0}, True),
        20: BasalGangliaSlice(0.0, {"Lentiform Nucleus": 35.0}, False),
    }
    context = StudyExportContext("CT", 40, 0.5, None, [], basal_ganglia=slices)
    export_study_csv(_study_results(), context, tmp_path / "s.csv", tmp_path / "p.csv")
    summary = _read_csv(tmp_path / "s.csv")
    assert ["13", "clearly visible", "0.900", "6.50", "7.00", "yes"] in summary
    assert ["21", "label not plausible", "0.000", "35.00", "", "no"] in summary
    per_slice = _read_csv(tmp_path / "p.csv")
    column = per_slice[0].index("Basal ganglia (provisional)")
    assert {row[0]: row[column] for row in per_slice[1:] if row[1] == "Falx"}["13"] == "clearly visible"

    volume = VolumeQa("S", "study", slice_count=40, segmented_slices=40,
                      structures={"Dural Folds": StructureQa("Dural Folds")}, basal_ganglia=slices)
    export_dataset_csv([volume], tmp_path / "ds.csv", tmp_path / "dp.csv")
    rows = _read_csv(tmp_path / "ds.csv")
    assert rows[1][rows[0].index("Basal ganglia visible slices (provisional)")] == "13"
    assert rows[1][rows[0].index("Basal ganglia label not plausible slices")] == "21"


def test_segmentation_check_rows():
    from bictool.domain.segmentation.segmentation_plausibility import SegmentationPlausibility
    from bictool.integration.results_export import _plausibility_rows

    rows = _plausibility_rows(SegmentationPlausibility(fragmentation=38.0, slices=28, worst_slices=(13, 18, 11), lobe_islands=14))
    assert ["Segmentation", "may be unreliable - check the overlays"] in rows
    assert ["Fragmentation (pieces per slice, median)", "38.0", "limit 30"] in rows
    assert ["Most fragmented slices", "14, 19, 12"] in rows
    assert _plausibility_rows(SegmentationPlausibility(12.0, 28, (1,), 0))[2] == ["Segmentation", "looks consistent"]
    assert _plausibility_rows(None) == [] and _plausibility_rows(SegmentationPlausibility(0.0, 0, (), 0)) == []


def test_the_metal_rule_is_exported_with_the_model_verdict(tmp_path):
    from dataclasses import replace

    from bictool.domain.results.head_ct_verdict import NOT_ASSESSABLE, apply_metal_rule

    results = _physical_results()
    verdict, context = _physical_context(results)
    ruled = apply_metal_rule(verdict, [verdict.slice_ids[0]])
    export_study_csv(results, replace(context, head_ct_verdict=ruled), tmp_path / "s.csv", tmp_path / "p.csv")
    summary = _read_csv(tmp_path / "s.csv")

    assert ["Verdict", NOT_ASSESSABLE] in summary
    rule = next(row for row in summary if row[:1] == ["Metal rule"])
    assert rule[1].endswith(f"slices {verdict.slice_ids[0] + 1}") and rule[2] == f"model alone: {verdict.verdict}"
